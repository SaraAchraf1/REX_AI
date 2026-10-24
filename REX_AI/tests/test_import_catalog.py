"""Catalogue import: PDF / Word / Excel / CSV are always MERGED into the course data REX recommends from."""
import json
import shutil
from pathlib import Path

import pytest

from rex_ai.core import catalog_import as ci

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_PDF = ROOT / "data" / "catalog" / "examples" / "sample_academy_catalog.pdf"
SAMPLE_IDS = {"AI-301", "DS-201", "SW-301", "CY-201", "BU-301"}


@pytest.fixture
def data_dir(tmp_path, monkeypatch):
    """A private copy of the sample data so the real files are never touched."""
    from rex_ai import config
    d = tmp_path / "data"
    shutil.copytree(ROOT / "tests" / "fixtures" / "sample_data", d)
    (d / "catalog").mkdir()
    monkeypatch.setattr(config, "DATA_DIR", d)
    yield d
    from rex_ai.agents.rex.retriever import RAGRetriever
    from rex_ai.core.rag import get_rag
    get_rag(RAGRetriever.COLLECTION).delete(where={"origin": RAGRetriever.DOCS_ORIGIN})  # keep other tests clean


def _courses(d):
    return {c["course_id"]: c for c in json.loads((d / "courses.json").read_text(encoding="utf-8"))["courses"]}


def _prices(d):
    return {p["course_id"]: p for p in json.loads((d / "pricing_promotions.json").read_text(encoding="utf-8"))["prices"]}


def test_sample_pdf_parses_five_valid_courses():
    res = ci.parse_file(SAMPLE_PDF, {"INS-01": "A", "INS-04": "D", "INS-06": "F", "INS-07": "G", "INS-09": "I"})
    assert [c["course_id"] for c in res.courses] == ["AI-301", "DS-201", "SW-301", "CY-201", "BU-301"]
    assert not res.errors
    ai = res.courses[0]
    assert ai["title_en"] == "Deep Learning with PyTorch"
    assert ai["formats"] == ["online_live", "hybrid"] and ai["country"] == ["EG", "SA"]
    assert ai["description_en"].endswith("deployed capstone.")             # wrapped lines joined
    assert res.prices[0] == {"course_id": "AI-301", "fee": 16500, "installments": True}
    assert [s["start_date"] for s in res.schedules if s["course_id"] == "AI-301"] == ["2026-12-06", "2027-01-17"]


def test_merge_keeps_old_and_adds_new(data_dir):
    old = set(_courses(data_dir))
    assert ci.main([str(SAMPLE_PDF)]) == 0
    now = _courses(data_dir)
    assert old | SAMPLE_IDS == set(now)                                     # nothing removed
    assert set(_prices(data_dir)) >= old | SAMPLE_IDS
    assert list((data_dir / "backup").iterdir())


def test_same_course_id_is_updated_not_duplicated(data_dir, tmp_path):
    ci.main([str(SAMPLE_PDF)])
    update = tmp_path / "update.txt"
    update.write_text("Course ID: AI-301\nTitle: Deep Learning with PyTorch (2027 edition)\nDepartment: AI\n"
                      "Level: Advanced\nDuration (hours): 60\nDescription: Updated.\nTags: ai; deep learning\n"
                      "Fee (EGP): 18,000\n", encoding="utf-8")
    report = ci.import_paths([update])
    assert report["merge"]["updated"] == ["AI-301"] and report["merge"]["added"] == []
    c = _courses(data_dir)
    assert list(c).count("AI-301") == 1 and c["AI-301"]["duration_hours"] == 60
    assert _prices(data_dir)["AI-301"]["fee"] == 18000
    assert "DS-101" in c                                                    # old sample course still there


def test_catalog_folder_all_file_types_and_unchanged_skipped(data_dir):
    import docx
    import pandas as pd
    folder = data_dir / "catalog"
    shutil.copy(SAMPLE_PDF, folder / "catalog.pdf")
    d = docx.Document()
    t = d.add_table(rows=0, cols=2)
    for k, v in [("Course ID", "ZZ-101"), ("Title", "Word Course"), ("Department", "AI"), ("Level", "Beginner"),
                 ("Duration (hours)", "10"), ("Description", "From Word."), ("Tags", "ai; machine learning")]:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
    d.save(folder / "word.docx")
    pd.DataFrame([{"Course ID": "ZZ-102", "Title": "Excel Course", "Department": "Data", "Level": "Intermediate",
                   "Duration (hours)": "20", "Description": "From Excel.", "Formats": "Online live; Onsite",
                   "Tags": "data analysis; sql", "Fee (EGP)": "5000"}]).to_excel(folder / "sheet.xlsx", index=False)
    (folder / "more.csv").write_text("Course ID,Title,Department,Level,Duration (hours),Description,Tags\n"
                                     "ZZ-103,CSV Course,Cloud,Beginner,12,From CSV.,cloud; aws\n", encoding="utf-8")

    report = ci.import_paths([folder])
    assert set(report["merge"]["added"]) == SAMPLE_IDS | {"ZZ-101", "ZZ-102", "ZZ-103"}
    assert _courses(data_dir)["ZZ-102"]["formats"] == ["online_live", "onsite"]
    assert "DS-101" in _courses(data_dir)

    again = ci.import_paths([folder])                                       # nothing changed
    assert all(r.skipped for r in again["files"]) and again["merge"] is None


def test_auto_import_on_api_startup(data_dir, monkeypatch):
    from fastapi.testclient import TestClient
    from rex_ai import config
    from rex_ai.agents.rex.knowledge_base import get_kb
    shutil.copy(SAMPLE_PDF, data_dir / "catalog" / "catalog.pdf")
    monkeypatch.setattr(config, "AUTO_IMPORT", True)
    get_kb.cache_clear()
    from rex_ai.api.main import app
    with TestClient(app) as client:
        courses = {c["course_id"] for c in client.get("/courses").json()}
    assert SAMPLE_IDS <= courses and "DS-101" in courses
    get_kb.cache_clear()


def test_dry_run_writes_nothing(data_dir):
    before = (data_dir / "courses.json").read_text(encoding="utf-8")
    assert ci.main([str(SAMPLE_PDF), "--dry-run"]) == 0
    assert (data_dir / "courses.json").read_text(encoding="utf-8") == before


def test_invalid_course_is_reported_not_imported(tmp_path, data_dir, capsys):
    bad = tmp_path / "bad.txt"
    bad.write_text("Course ID: BAD-1\nTitle: No level or hours\nDepartment: AI\nDescription: x\n", encoding="utf-8")
    assert ci.main([str(bad)]) == 1
    assert "BAD-1" not in _courses(data_dir)
    assert "missing required field 'skill_level'" in capsys.readouterr().out

def _guide(path, text):
    import docx
    d = docx.Document()
    d.add_heading("Campus guide", 1)
    d.add_paragraph(text)
    d.save(path)
    return path.read_bytes()


def test_uploaded_document_is_answered_and_never_overwritten(data_dir, tmp_path):
    from rex_ai.core import data_sync
    from rex_ai.core.schemas import ChatRequest
    first = _guide(tmp_path / "a.docx", "Parking: free parking for learners is available in the basement of the "
                                        "Nasr City campus from 9 am to 10 pm. Show your learner ID card at the gate.")
    res = data_sync.save_upload("guide.docx", first)
    assert res["status"] == "saved" and res["kind"] == "document" and res["file"] == "documents/guide.docx"
    assert data_sync.save_upload("guide.docx", first)["status"] == "duplicate"
    second = _guide(tmp_path / "b.docx", "Library: the academy library lends laptops for the length of a course.")
    assert data_sync.save_upload("guide.docx", second)["file"] == "documents/guide-2.docx"
    assert (data_dir / "documents" / "guide.docx").read_bytes() == first          # old file untouched

    rex = data_sync.get_assistant()
    r = rex.handle(ChatRequest(session_id="docs-1", message="Where is the free parking at the campus?"))
    assert r.intent == "documents" and "basement" in r.reply
    assert r.sources[0].source_id == "SRC-ACAD-DOCS"
    assert "DS-101" in rex.kb.courses


def test_new_catalogue_picked_up_without_restart(data_dir):
    from rex_ai.core import data_sync
    before = data_sync.get_assistant()
    assert not SAMPLE_IDS & set(before.kb.courses)
    (data_dir / "catalog" / "examples").mkdir()
    shutil.copy(SAMPLE_PDF, data_dir / "catalog" / "examples" / "catalog.pdf")   # sub-folders are imported
    after = data_sync.get_assistant()
    assert after is not before and SAMPLE_IDS | set(before.kb.courses) == set(after.kb.courses)


def test_api_adds_courses_and_uploads(data_dir):
    from fastapi.testclient import TestClient
    from rex_ai.api.main import app
    with TestClient(app) as client:
        r = client.post("/data/courses", json={"courses": [{
            "course_id": "AI-401", "title_en": "MLOps in Production", "department": "AI", "skill_level": "Advanced",
            "duration_hours": 40, "description_en": "Ship and monitor ML models.", "tags": ["mlops", "machine learning"],
            "schedule": [{"start_date": "2099-01-10", "format": "Online live", "days": "Sat", "time": "19:00-22:00"}]}]})
        assert r.status_code == 200 and r.json()["added"] == ["AI-401"]
        reply = client.post("/chat", json={"session_id": "api-data", "message": "tell me about AI-401"}).json()
        assert "MLOps" in reply["reply"] and "2099-01-10" in reply["reply"]

        csv = b"Course ID,Title,Department,Level,Duration (hours),Description,Tags,Fee (EGP)\n" \
              b"ZZ-201,CSV Upload Course,Cloud,Beginner,12,From an upload.,cloud; aws,4000\n"
        up = client.post("/data/upload", files={"file": ("new.csv", csv)}).json()
        assert up["kind"] == "catalog" and up["added"] == ["ZZ-201"]
        assert client.post("/data/upload", files={"file": ("bad.exe", b"x")}).status_code == 400
        ids = {c["course_id"] for c in client.get("/courses").json()}
        assert {"AI-401", "ZZ-201", "DS-101"} <= ids
        assert "catalog/new.csv" in client.get("/data/status").json()["catalog_files"]


def test_unanswerable_question_says_so(chat):
    for q in ["How do I become a pilot?", "What is the wifi password?", "هل يوجد مطعم في الأكاديمية؟"]:
        r = chat(q)
        assert r.intent in ("unknown", "documents"), (q, r.intent)
        if r.intent == "unknown":
            assert not r.recommendations and not r.sources
    assert chat("I want to learn data analysis").intent == "discovery"       # profile info still starts discovery
    assert chat("Which course is best for a beginner in AI?").intent != "unknown"


def test_line_endings_do_not_count_as_a_changed_catalogue(tmp_path):
    """A Windows checkout (core.autocrlf) must not trigger a re-import of an unchanged catalogue."""
    lf, crlf = tmp_path / "lf.txt", tmp_path / "crlf.txt"
    lf.write_bytes(b"Course ID: AI-101\nTitle: Intro\n")
    crlf.write_bytes(b"Course ID: AI-101\r\nTitle: Intro\r\n")
    assert ci._file_hash(lf) == ci._file_hash(crlf)
    assert json.loads((ROOT / "data" / "catalog" / ci.REGISTRY).read_text(encoding="utf-8"))[
        "examples/sample_academy_catalog.pdf"] == ci._file_hash(SAMPLE_PDF)
