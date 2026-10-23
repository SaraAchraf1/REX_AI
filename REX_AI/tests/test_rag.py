"""Shared RAG pipeline tests (offline hashing embedder, no model download).

PDF tests run on the REAL MIT OCW 6.7960 lecture notes downloaded by scripts/ingest_mit_ocw.py.
Download them once (only lectures 5 and 8 are needed for the tests):

    python scripts/ingest_mit_ocw.py --lectures 5 8

If the PDFs are not there, the PDF tests are skipped and the Word / Excel / chunking tests still run.
Set REX_TEST_PDF_DIR to use another folder.
"""
import os
import re
import sys
from pathlib import Path

import pytest

from rex_ai.core.rag import (ChunkingConfig, HashingEmbedder, RAGPipeline, load_document, recursive_split,
                             semantic_split)

ROOT = Path(__file__).resolve().parents[1]
PDF_DIR = Path(os.getenv("REX_TEST_PDF_DIR", ROOT / "data" / "raw" / "6-7960-deep-learning-fall-2024"))
LEC5 = PDF_DIR / "mit6_7960_f24_lec5.pdf"     # Graph Neural Networks (has embedded LaTeX equations)
LEC8 = PDF_DIR / "mit6_7960_f24_lec8.pdf"     # Transformers

needs_pdfs = pytest.mark.skipif(
    not (LEC5.exists() and LEC8.exists()),
    reason=f"real MIT PDFs not found in {PDF_DIR} - run: python scripts/ingest_mit_ocw.py --lectures 5 8")

NOISE = re.compile(r"latexit|sha1_base64|\(cid:\d+\)|[A-Za-z0-9+/]{40,}", re.I)


# =============================================================================
# Fixtures
# =============================================================================
@pytest.fixture(scope="module")
def office_files(tmp_path_factory):
    """Small Word and Excel files (python-docx / pandas are already project dependencies)."""
    d = tmp_path_factory.mktemp("docs")
    import docx
    doc = docx.Document()
    doc.add_heading("Data Analysis Course Proposal", 0)
    doc.add_heading("Learning Outcomes", 1)
    doc.add_paragraph("Clean data with pandas and build Power BI dashboards.")
    doc.add_heading("Fees", 1)
    t = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Format", "Fee (EGP)"), ("Online", "10500"), ("Onsite", "12500")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    doc.add_heading("الوصف", 1)
    doc.add_paragraph("دورة تحليل البيانات باستخدام بايثون للمبتدئين.")
    word = d / "proposal.docx"
    doc.save(word)

    import pandas as pd
    xlsx = d / "catalog.xlsx"
    with pd.ExcelWriter(xlsx) as w:
        pd.DataFrame({"course_id": ["AI-201", "DS-101", "CY-110"],
                      "title": ["Deep Learning Bootcamp", "Data Analysis with Python", "Cybersecurity Essentials"],
                      "level": ["Advanced", "Beginner", "Beginner"]}).to_excel(w, sheet_name="Courses", index=False)
        pd.DataFrame({"instructor": ["Dr. Salma"], "expertise": ["Computer vision"]}).to_excel(
            w, sheet_name="Instructors", index=False)
    return {"docx": word, "xlsx": xlsx}


@pytest.fixture(scope="module")
def rag(tmp_path_factory, office_files):
    r = RAGPipeline("test_office", persist_dir=tmp_path_factory.mktemp("chroma"), embedder=HashingEmbedder(),
                    chunking=ChunkingConfig(chunk_size=400, chunk_overlap=60, min_chunk_size=80))
    report = r.ingest(list(office_files.values()), metadata={"source_id": "SRC-TEST", "department": "AI"})
    assert report["errors"] == [] and report["files"] == 2
    return r


@pytest.fixture(scope="module")
def mit_rag(tmp_path_factory):
    """The two real lectures, ingested the same way scripts/ingest_mit_ocw.py does it."""
    r = RAGPipeline("test_mit", persist_dir=tmp_path_factory.mktemp("chroma_mit"), embedder=HashingEmbedder())
    for pdf, lec, title in [(LEC5, 5, "Lecture 5: Graph Neural Networks"), (LEC8, 8, "Lecture 8: Transformers")]:
        assert r.ingest_file(pdf, metadata={"title": title, "lecture": lec, "source_id": "SRC-MIT-OCW-6.7960",
                                            "department": "AI"}) > 0
    return r


# =============================================================================
# Real MIT lecture PDFs
# =============================================================================
@needs_pdfs
def test_pdf_loader_keeps_pages_and_slide_titles():
    doc = load_document(LEC5)
    assert len(doc.blocks) > 10
    assert all(b.page for b in doc.blocks)
    assert sum(1 for b in doc.blocks if b.heading_path) > len(doc.blocks) / 2      # most text sits under a title


@needs_pdfs
def test_pdf_embedded_latex_is_removed(mit_rag):
    """The slides hide equation source as <latexit sha1_base64=...>AAAC...</latexit>; none may reach a chunk."""
    docs = mit_rag.collection.get(include=["documents"])["documents"]
    bad = [d[:120] for d in docs if NOISE.search(d)]
    assert not bad, f"{len(bad)} chunks still contain LaTeX/base64 noise, e.g. {bad[:2]}"


@needs_pdfs
def test_pdf_headers_have_no_duplicate_title_or_citation_sections(mit_rag):
    data = mit_rag.collection.get(include=["documents", "metadatas"])
    for text, meta in zip(data["documents"], data["metadatas"]):
        assert text.startswith(f"[{meta['title']}")
        assert not text.startswith(f"[{meta['title']} > {meta['title']}")
        assert "Lecture 5" not in meta.get("section", "") and "Lecture 8" not in meta.get("section", "")
        assert not re.search(r"\bet al\b", meta.get("section", ""))               # citations are not headings


@needs_pdfs
def test_search_returns_the_right_lecture_with_page(mit_rag):
    gnn = mit_rag.search("graph neural network message passing between nodes", k=3)
    attn = mit_rag.search("transformer self-attention queries keys values", k=3)
    assert 5 in [h.metadata["lecture"] for h in gnn]
    assert 8 in [h.metadata["lecture"] for h in attn]
    assert "p." in gnn[0].citation


@needs_pdfs
def test_pdf_filter_and_source_refs(mit_rag):
    hits = mit_rag.search("attention", k=3, where={"lecture": 8})
    assert hits and all(h.metadata["lecture"] == 8 for h in hits)
    ref = mit_rag.to_source_refs(hits)[0]
    assert ref.source_id == "SRC-MIT-OCW-6.7960" and ref.record_id.startswith("mit6_7960_f24_lec8.pdf#page=")


@needs_pdfs
def test_unchanged_pdf_is_skipped(mit_rag):
    before = mit_rag.collection.count()
    assert mit_rag.ingest_file(LEC8, metadata={"title": "Lecture 8: Transformers"}) == 0
    assert mit_rag.collection.count() == before


# =============================================================================
# Word / Excel
# =============================================================================
def test_docx_loader_keeps_headings_and_tables(office_files):
    doc = load_document(office_files["docx"])
    table = next(b for b in doc.blocks if b.kind == "table")
    assert "Format: Online; Fee (EGP): 10500" in table.text
    assert table.heading_path == ["Data Analysis Course Proposal", "Fees"]
    assert any("الوصف" in b.heading_path for b in doc.blocks)


def test_excel_loader_one_block_per_row_per_sheet(office_files):
    doc = load_document(office_files["xlsx"])
    assert {b.extra["sheet"] for b in doc.blocks} == {"Courses", "Instructors"}
    assert doc.blocks[0].text.startswith("course_id: AI-201; title: Deep Learning Bootcamp")
    assert doc.blocks[0].extra["row"] == 2


def test_chunks_never_cross_sections_and_carry_context(rag):
    data = rag.collection.get(where={"source": "proposal.docx"}, include=["documents", "metadatas"])
    for text, meta in zip(data["documents"], data["metadatas"]):
        assert text.startswith("[proposal")
        if meta.get("section", "").endswith("Fees"):
            assert "10500" in text and "pandas" not in text
        if meta.get("section", "").endswith("Learning Outcomes"):
            assert "pandas" in text and "10500" not in text


def test_search_metadata_filter_and_excel(rag):
    hits = rag.search("beginner data analysis course", k=3, where={"file_type": "xlsx"})
    assert hits and all(h.metadata["file_type"] == "xlsx" for h in hits)
    assert "DS-101" in hits[0].text and hits[0].metadata["sheet"] == "Courses"


def test_arabic_query(rag):
    hits = rag.search("تحليل البيانات بايثون", k=2)
    assert any("الوصف" in h.metadata.get("section", "") for h in hits)


def test_reingest_unchanged_file_is_skipped(rag, office_files):
    before = rag.collection.count()
    assert rag.ingest_file(office_files["xlsx"]) == 0
    assert rag.collection.count() == before


def test_older_pipeline_version_is_reingested(tmp_path, office_files):
    r = RAGPipeline("versions", persist_dir=tmp_path, embedder=HashingEmbedder())
    n = r.ingest_file(office_files["docx"])
    ids = r.collection.get()["ids"]
    r.collection.update(ids=ids, metadatas=[{"pipeline_version": "1"}] * len(ids))   # simulate an old ingest
    assert r.ingest_file(office_files["docx"]) == n                                 # re-ingested, not skipped
    assert r.collection.count() == n                                                # old chunks replaced


def test_context_and_source_refs(rag):
    hits = rag.search("course fees online onsite", k=2)
    assert rag.build_context(hits).startswith("[1] Source: proposal")
    assert rag.to_source_refs(hits)[0].source_id == "SRC-TEST"


def test_add_texts_for_existing_kb_records(tmp_path):
    r = RAGPipeline("records", persist_dir=tmp_path, embedder=HashingEmbedder())
    r.add_texts(["Deep Learning Bootcamp for engineers", "Accounting basics"],
                [{"type": "course", "course_id": "AI-201"}, {"type": "course", "course_id": "AC-100"}],
                ids=["AI-201#overview", "AC-100#overview"])
    assert r.search("deep learning", k=1, where={"type": "course"})[0].metadata["course_id"] == "AI-201"


# =============================================================================
# Chunking + noise helpers
# =============================================================================
def test_recursive_split_respects_size_and_overlap():
    text = " ".join(f"Sentence number {i} talks about neural networks." for i in range(60))
    parts = recursive_split(text, 200, 40)
    assert len(parts) > 5 and all(len(p) <= 200 for p in parts)
    assert parts[0][-20:].split()[-1] in parts[1]          # overlap carried into the next chunk


def test_semantic_split_cuts_at_topic_change():
    text = ("Cats are small furry pets. Cats purr and sleep a lot. Kittens are baby cats. Cats chase mice. "
            "Quarterly revenue grew by ten percent. The finance team closed the books. Revenue and profit rose. "
            "Investors liked the earnings report.")
    groups = semantic_split(text, HashingEmbedder(), percentile=80, min_size=50)
    assert len(groups) >= 2
    assert "Cats" in groups[0] and "revenue" in groups[-1].lower()


def test_latexit_text_is_stripped():
    from rex_ai.core.rag import _clean
    raw = ('Message passing updates each node. <latexit sha1_base64="YJe67Lcoy1dxBPddIP+o/htLYsE=">'
           'AAACcnicbVHdShwxGM2MttVttVuLNxVsdBEUZJkpovVCEOqFV2Khq8LOMG</latexit> Permutation invariance.')
    cleaned = _clean(raw)
    assert not NOISE.search(cleaned)
    assert "Message passing" in cleaned and "Permutation invariance" in cleaned


# =============================================================================
# OCW scraper (no network: the list page is mocked)
# =============================================================================
def test_ocw_scraper_parses_list_page(monkeypatch):
    sys.path.insert(0, str(ROOT / "scripts"))
    import ingest_mit_ocw as ocw

    html = """<a href="/courses/6-7960-deep-learning-fall-2024/mit6_7960_f24_lec8.pdf">pdf 13 MB</a>
              <a href="/courses/6-7960-deep-learning-fall-2024/resources/mit6_7960_f24_lec8_pdf/">Transformers</a>
              <a href="/courses/6-7960-deep-learning-fall-2024/resources/mit6_7960_f24_lec5_pdf/">Graph Neural Networks</a>"""
    resource = '<a href="/courses/6-7960-deep-learning-fall-2024/mit6_7960_f24_lec5.pdf">Download file</a>'

    class R:
        def __init__(self, t):
            self.text = t

        def raise_for_status(self):
            pass

    monkeypatch.setattr(ocw.SESSION, "get", lambda url, **kw: R(resource if "/resources/" in url else html))
    lectures = ocw.find_lectures(ocw.DEFAULT_URL)
    assert [(lec["lecture"], lec["title"]) for lec in lectures] == [(5, "Graph Neural Networks"), (8, "Transformers")]
    assert lectures[1]["pdf_url"] == "https://ocw.mit.edu/courses/6-7960-deep-learning-fall-2024/mit6_7960_f24_lec8.pdf"
    assert lectures[0]["pdf_url"].endswith("lec5.pdf")


@pytest.mark.parametrize("raw, expected", [
    ('w: <latexit sha1_base64="f', "w:"),                                  # fragment split by the PDF reader
    ('sha1_base64="flGO7PvVKXBCQ', ""),
    ("AAAB8XicbVBNS0NBEJ2 hello world", "hello world"),                 # short base64 fragment
    ("PyTorchLightning input/output ResNet50 GPT-4o CIFAR10 ImageNet1k",   # real words are kept
     "PyTorchLightning input/output ResNet50 GPT-4o CIFAR10 ImageNet1k"),
])
def test_noise_fragments(raw, expected):
    from rex_ai.core.rag import _clean
    assert _clean(raw) == expected


def test_layout_artefact_tables_are_ignored():
    from rex_ai.core.rag import _render_table
    assert _render_table([["General strategy"], ["this,"], [None]]) == ""             # one column = not a table
    assert _render_table([["w", None, None, None, None], ["", "x", "", "", ""], ["", "", "", "", "y"]]) == ""
    assert _render_table([["Model", "Params"], ["GPT-2", "1.5B"], ["BERT", "340M"]]) == \
        "Model: GPT-2; Params: 1.5B\nModel: BERT; Params: 340M"


def test_ocw_falls_back_to_builtin_list_when_offline(monkeypatch):
    sys.path.insert(0, str(ROOT / "scripts"))
    import ingest_mit_ocw as ocw
    lectures = ocw.known_lectures()
    assert len(lectures) == 23 and 22 not in [lec["lecture"] for lec in lectures]
    lec8 = next(lec for lec in lectures if lec["lecture"] == 8)
    assert lec8["title"] == "Transformers" and lec8["pdf_url"].endswith("/mit6_7960_f24_lec8.pdf")


def test_cid_placeholders_removed():
    from rex_ai.core.rag import _cid_ratio, _clean
    raw = "(cid:55)hat(cid:189)s the stride"
    assert _cid_ratio(raw) > 0.05 and "(cid:" not in _clean(raw)


def test_numbered_bullets_are_not_headings_when_pdf_has_title_font():
    from rex_ai.core.rag import _is_pdf_heading
    titles = [36.0]
    assert _is_pdf_heading("Two goals", 36.0, 24.0, 4, titles)
    assert not _is_pdf_heading("1. Node embeddings", 24.0, 24.0, 4, titles)
    assert not _is_pdf_heading("1. Encode each node (based on message passing between nodes) 2. Aggregate set",
                               36.0, 24.0, 4, titles)                       # too long for a title


def test_small_slides_are_merged_and_keep_their_titles():
    from rex_ai.core.rag import Block, Chunker, LoadedDocument
    blocks = [Block("1. Node embeddings\n2. Graph embedding", ["Two goals"], page=1),
              Block("Encode each node with message passing.", ["Message passing"], page=2),
              Block("Node order must not change the output.", ["Permutation invariance"], page=3)]
    doc = LoadedDocument("d1", "lec5.pdf", "Lecture 5: Graph Neural Networks", "pdf", blocks)
    chunks = Chunker(ChunkingConfig(chunk_size=400, min_chunk_size=200)).chunk(doc)
    assert len(chunks) == 1
    c = chunks[0]
    assert c.text.startswith("[Lecture 5: Graph Neural Networks > Two goals]")
    assert "Message passing:" in c.text and "Permutation invariance:" in c.text
    assert c.metadata["page_start"] == 1 and c.metadata["page_end"] == 3
    assert c.metadata["also_covers"] == "Message passing | Permutation invariance"


def test_garbled_font_text_is_detected():
    from rex_ai.core.rag import _is_garbled
    assert _is_garbled("U earning tass with gra«hs U essage «assing NNs U Ƃ««roximation *ower")
    assert _is_garbled("(cid:55)hat(cid:189)s the (cid:142)ernel si(cid:226)e")
    assert not _is_garbled("A ∈ ℝ n×n X ∈ ℝ n×d Two goals 1. Node embeddings 2. Graph embedding")
    assert not _is_garbled("دورة تحليل البيانات باستخدام بايثون للمبتدئين")


def test_slide_titles_from_largest_top_line():
    from rex_ai.core.rag import _pdf_blocks_slides

    def page(lines, source="text"):
        return {"lines": [{"text": t, "size": sz, "top": top} for t, sz, top in lines], "tables": [],
                "source": source, "landscape": True, "height": 540.0}

    blocks = _pdf_blocks_slides([
        page([("Two goals", 36, 40), ("1. Node embeddings", 24, 150), ("2. Graph embedding", 24, 200)]),
        page([("Message passing", 36, 40), ("1. Encode each node (based on message passing between nodes)", 24, 150)]),
        page([("pooling operation (just like AGGREGATE)", 24, 150), ("12", 10, 520)]),   # continuation, no title
    ])
    assert [b.heading_path for b in blocks] == [["Two goals"], ["Message passing"], ["Message passing"]]
    assert blocks[0].text == "1. Node embeddings\n2. Graph embedding"
    assert "12" not in blocks[2].text


@needs_pdfs
def test_pdf_chunks_are_not_garbled(mit_rag):
    from rex_ai.core.rag import _is_garbled
    docs = mit_rag.collection.get(include=["documents"])["documents"]
    bad = [d[:150] for d in docs if _is_garbled(d)]
    assert not bad, f"{len(bad)} garbled chunks (install OCR: pip install rapidocr_onnxruntime), e.g. {bad[:2]}"


def test_pdf_helper_functions_exist():
    """Guards against a helper being dropped by a refactor (every fallback path must be callable)."""
    from rex_ai.core import rag
    for name in ("_page_text_pdfium", "_ocr_page", "_words_to_lines", "_pdf_blocks_slides",
                 "_pdf_blocks_document", "_is_garbled", "_inside", "_is_pdf_heading", "_render_table"):
        assert callable(getattr(rag, name)), name
    assert rag._page_text_pdfium(Path("does-not-exist.pdf"), 0) is None


def test_garbled_page_falls_back_to_pdfium(tmp_path, monkeypatch):
    """Force the 'garbled font' branch of load_pdf on a tiny hand-written PDF and check pdfium is used."""
    from rex_ai.core import rag
    pdf = tmp_path / "mini.pdf"
    pdf.write_bytes(_minimal_pdf("Graph neural networks pass messages between nodes."))
    monkeypatch.setattr(rag, "_is_garbled", lambda text: "pdfium" not in calls)
    calls = []
    real = rag._page_text_pdfium

    def spy(path, index):
        calls.append("pdfium")
        return real(path, index)

    monkeypatch.setattr(rag, "_page_text_pdfium", spy)
    blocks = rag.load_pdf(pdf)
    assert calls == ["pdfium"]
    assert blocks and "messages between nodes" in blocks[0].text
    assert blocks[0].extra["text_source"] == "pdfium"


def _minimal_pdf(text: str) -> bytes:
    """A valid one-page PDF with one line of Helvetica text (no extra libraries needed)."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


@pytest.mark.parametrize("line, noise", [
    ("oFαTo■oTαTo", True), ("{ 0o 0", True), ("UC", True), ("C", True), ("olαTo■olTo", True),
    ("A CNN is a GNN over a grid graph", False), ("· What's the kernel size?", False),
    ("Self-attention: softmax(QK^T/sqrt(d))V", False), ("α-β pruning in search", False),
])
def test_ocr_figure_noise_filter(line, noise):
    from rex_ai.core.rag import _is_figure_noise
    assert _is_figure_noise(line) is noise